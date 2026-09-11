import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.core.exceptions import TimetablePublishError
from app.models.enums import TimetableRunStatus
from app.models.time_slot import TimeSlot
from app.models.timetable_entry import TimetableEntry
from app.services.timetable.timetable_run import get_timetable_run
from app.services.timetable.validation import validate_timetable_run
from sqlalchemy import select
from sqlalchemy.orm import Session


XLSX_HEADER = [
    "Time",
    "Monday",
    "Tuesday",
    "Wednesday",
    "Thursday",
    "Friday",
]


DARK_BLUE = "1F4E78"
MEDIUM_BLUE = "D9EAF7"
LIGHT_BLUE = "F4F8FB"
TEXT_COLOR = "1F2937"
BORDER_COLOR = "D9E2F3"

THIN_BORDER = Side(
    style="thin",
    color=BORDER_COLOR,
)


def _group_entries_by_faculty_program_level_and_semester(
    entries: tuple[TimetableEntry, ...],
) -> list[tuple[tuple[str, str, str, str], list[TimetableEntry]]]:
    sections: dict[tuple[str, str, str, str], list[TimetableEntry]] = {}

    for entry in entries:
        curriculum = entry.course_session.course_offering.curriculum_course
        program = curriculum.program_semester.study_program
        level = program.level
        faculty = program.faculty
        semester = curriculum.program_semester

        section_key = (
            f"{faculty.code} — {faculty.name}",
            f"{program.code} — {program.name}",
            f"{level.code} — {level.name}",
            f"Semester {semester.semester_number}",
        )

        sections.setdefault(section_key, []).append(entry)

    return sorted(
        sections.items(),
        key=lambda item: (
            item[0][0].casefold(),
            item[0][1].casefold(),
            item[0][2].casefold(),
            int(item[0][3].split()[-1]),
        ),
    )


def _entry_cell_text(
    entry: TimetableEntry,
    slot_by_position: dict[tuple[int, int], TimeSlot],
) -> str:
    session = entry.course_session
    curriculum = session.course_offering.curriculum_course
    course = curriculum.course

    start_slot = entry.start_slot

    end_slot = slot_by_position[
        (
            start_slot.day_of_week,
            start_slot.slot_index + session.duration_slots - 1,
        )
    ]

    groups = ", ".join(
        assignment.student_group.name
        for assignment in session.group_assignments
    )

    staff = ", ".join(
        f"{assignment.staff_member.first_name} "
        f"{assignment.staff_member.last_name}"
        for assignment in session.staff_assignments
    )

    start_time = start_slot.start_time.isoformat(
        timespec="minutes"
    )
    end_time = end_slot.end_time.isoformat(
        timespec="minutes"
    )

    details = [
        f"{course.code} — {course.name}",
        (
            f"{session.name} · "
            f"{session.component_type.value.title()} · "
            f"{start_time}–{end_time}"
        ),
        f"Room {entry.room.code}",
    ]

    if groups:
        details.append(f"Groups: {groups}")

    if staff:
        details.append(f"Staff: {staff}")

    return "\n".join(details)


def _sheet_title(
    faculty_label: str,
    program_label: str,
    level_label: str,
    semester_label: str,
    used_titles: set[str],
) -> str:
    parts = [
        faculty_label.split(" — ", maxsplit=1)[0],
        program_label.split(" — ", maxsplit=1)[0],
        level_label.split(" — ", maxsplit=1)[0],
        semester_label.replace("Semester ", "S"),
    ]

    base = re.sub(
        r"[\\/*?:\[\]]",
        "",
        " ".join(parts),
    )[:31].strip()

    if not base:
        base = "Timetable"

    title = base
    suffix = 2

    while title in used_titles:
        suffix_text = f" {suffix}"
        title = (
            f"{base[:31 - len(suffix_text)]}"
            f"{suffix_text}"
        )
        suffix += 1

    used_titles.add(title)
    return title


def _style_timetable_sheet(
    worksheet,
    *,
    faculty_label: str,
    program_label: str,
    level_label: str,
    semester_label: str,
) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A6"

    metadata = [
        f"Faculty: {faculty_label}",
        f"Department / program: {program_label}",
        f"Study level: {level_label}",
        semester_label,
    ]

    for row_number, value in enumerate(metadata, start=1):
        worksheet.merge_cells(
            start_row=row_number,
            start_column=1,
            end_row=row_number,
            end_column=6,
        )

        cell = worksheet.cell(
            row=row_number,
            column=1,
            value=value,
        )

        cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
        )

        cell.font = Font(
            name="Arial",
            size=15 if row_number == 1 else 11,
            bold=row_number in (1, 4),
            color="FFFFFF" if row_number == 1 else DARK_BLUE,
        )

        cell.fill = PatternFill(
            "solid",
            fgColor=(
                DARK_BLUE
                if row_number == 1
                else MEDIUM_BLUE
                if row_number == 4
                else LIGHT_BLUE
            ),
        )

    worksheet.row_dimensions[1].height = 30
    worksheet.row_dimensions[2].height = 21
    worksheet.row_dimensions[3].height = 21
    worksheet.row_dimensions[4].height = 23

    worksheet.column_dimensions["A"].width = 14

    for column in "BCDEF":
        worksheet.column_dimensions[column].width = 42


def export_timetable_xlsx(
    db: Session,
    timetable_run_id: int,
) -> tuple[str, bytes]:
    run = get_timetable_run(
        db,
        timetable_run_id,
    )

    if run.status != TimetableRunStatus.SUCCEEDED:
        raise TimetablePublishError(
            "Only a successfully generated timetable can be downloaded."
        )

    validation = validate_timetable_run(
        db,
        run.id,
    )

    if not validation.is_valid:
        raise TimetablePublishError(
            "A timetable with hard conflicts cannot be downloaded.",
            details={
                "timetable_run_id": run.id,
                "hard_conflicts": validation.hard_conflict_count,
            },
        )

    entries = tuple(
        db.scalars(
            select(TimetableEntry)
            .join(
                TimeSlot,
                TimeSlot.id == TimetableEntry.start_slot_id,
            )
            .where(
                TimetableEntry.timetable_run_id == run.id
            )
            .order_by(
                TimeSlot.day_of_week,
                TimeSlot.slot_index,
                TimetableEntry.id,
            )
        ).all()
    )

    slots = tuple(
        db.scalars(
            select(TimeSlot).where(
                TimeSlot.scheduling_profile_id
                == run.scheduling_profile_id
            )
        ).all()
    )

    slot_by_position = {
        (
            slot.day_of_week,
            slot.slot_index,
        ): slot
        for slot in slots
    }

    workbook = Workbook()
    workbook.remove(workbook.active)

    workbook.properties.title = (
        f"Timetable run {run.id}"
    )

    used_sheet_titles: set[str] = set()

    sections = (
        _group_entries_by_faculty_program_level_and_semester(
            entries
        )
    )

    for (
        faculty_label,
        program_label,
        level_label,
        semester_label,
    ), section_entries in sections:
        worksheet = workbook.create_sheet(
            _sheet_title(
                faculty_label,
                program_label,
                level_label,
                semester_label,
                used_sheet_titles,
            )
        )

        _style_timetable_sheet(
            worksheet,
            faculty_label=faculty_label,
            program_label=program_label,
            level_label=level_label,
            semester_label=semester_label,
        )

        for column_index, header in enumerate(
            XLSX_HEADER,
            start=1,
        ):
            cell = worksheet.cell(
                row=5,
                column=column_index,
                value=header,
            )

            cell.font = Font(
                name="Arial",
                size=10,
                bold=True,
                color="FFFFFF",
            )

            cell.fill = PatternFill(
                "solid",
                fgColor=DARK_BLUE,
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

            cell.border = Border(
                bottom=Side(
                    style="medium",
                    color="FFFFFF",
                )
            )

        worksheet.row_dimensions[5].height = 26

        entries_by_day_and_start: dict[
            tuple[int, str],
            list[TimetableEntry],
        ] = {}

        start_times: set[str] = set()

        for entry in section_entries:
            start_time = entry.start_slot.start_time.isoformat(
                timespec="minutes"
            )

            start_times.add(start_time)

            entries_by_day_and_start.setdefault(
                (
                    entry.start_slot.day_of_week,
                    start_time,
                ),
                [],
            ).append(entry)

        for row_number, start_time in enumerate(
            sorted(start_times),
            start=6,
        ):
            row_values = [start_time]

            for day in range(1, 6):
                cell_entries = entries_by_day_and_start.get(
                    (
                        day,
                        start_time,
                    ),
                    [],
                )

                cell_text = "\n\n".join(
                    _entry_cell_text(
                        entry,
                        slot_by_position,
                    )
                    for entry in sorted(
                        cell_entries,
                        key=lambda item: (
                            item.start_slot.slot_index,
                            item.id,
                        ),
                    )
                )

                row_values.append(cell_text)

            for column_index, value in enumerate(
                row_values,
                start=1,
            ):
                cell = worksheet.cell(
                    row=row_number,
                    column=column_index,
                    value=value,
                )

                cell.font = Font(
                    name="Arial",
                    size=10,
                    color=TEXT_COLOR,
                )

                cell.alignment = Alignment(
                    horizontal="left",
                    vertical="top",
                    wrap_text=True,
                )

                cell.border = Border(
                    left=THIN_BORDER,
                    right=THIN_BORDER,
                    top=THIN_BORDER,
                    bottom=THIN_BORDER,
                )

                if column_index == 1:
                    cell.font = Font(
                        name="Arial",
                        size=10,
                        bold=True,
                        color=DARK_BLUE,
                    )

                    cell.fill = PatternFill(
                        "solid",
                        fgColor=MEDIUM_BLUE,
                    )

                    cell.alignment = Alignment(
                        horizontal="center",
                        vertical="top",
                    )

                elif row_number % 2 == 0:
                    cell.fill = PatternFill(
                        "solid",
                        fgColor=LIGHT_BLUE,
                    )

            worksheet.row_dimensions[row_number].height = 96

        last_row = max(
            5,
            5 + len(start_times),
        )

        worksheet.auto_filter.ref = (
            f"A5:F{last_row}"
        )

    if not workbook.worksheets:
        worksheet = workbook.create_sheet(
            "Timetable"
        )
        worksheet["A1"] = "No timetable entries"

    filename = f"timetable-run-{run.id}.xlsx"

    output = BytesIO()
    workbook.save(output)

    return filename, output.getvalue()